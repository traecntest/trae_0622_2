from __future__ import annotations

import copy
import re
from dataclasses import dataclass, field, asdict
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import config


@dataclass
class HeadingInfo:
    level: int
    text: str
    source_file: str
    para_index: int

    def to_dict(self) -> dict:
        return asdict(self)


@dataclass
class MergeResult:
    ok: bool
    output_path: str
    heading_count: int
    sources: list[str] = field(default_factory=list)
    toc: list[HeadingInfo] = field(default_factory=list)
    style_conflicts_resolved: int = 0
    error: str = ""

    def to_dict(self) -> dict:
        return {
            "ok": self.ok,
            "output_path": self.output_path,
            "heading_count": self.heading_count,
            "sources": self.sources,
            "toc": [h.to_dict() for h in self.toc],
            "style_conflicts_resolved": self.style_conflicts_resolved,
            "error": self.error,
        }


HEADING_RE = re.compile(r"^(第[一二三四五六七八九十百千]+章|第\d+章|[一二三四五六七八九十]+、|[（(]\d+[)）]|\d+[\.、])")
NUM_HEADING_RE = re.compile(r"^(\d+(\.\d+)*)\s*")


class DocMerger:
    def __init__(self, cfg: config.GWFormatConfig | None = None):
        self.cfg = cfg or config.GWFormatConfig()

    def _detect_heading_level(self, para) -> int:
        style_name = (para.style.name or "").lower() if para.style else ""
        text = (para.text or "").strip()
        if not text:
            return 0
        if "heading 1" in style_name or "title" in style_name:
            return 1
        if "heading 2" in style_name:
            return 2
        if "heading 3" in style_name:
            return 3
        if NUM_HEADING_RE.match(text):
            nums = NUM_HEADING_RE.match(text).group(1)
            level = nums.count(".") + 1
            return min(level, 4)
        if re.match(r"^(第[一二三四五六七八九十百千]+章)", text):
            return 1
        if re.match(r"^(第[一二三四五六七八九十百千]+节)", text):
            return 2
        if re.match(r"^[一二三四五六七八九十]+、", text):
            return 2
        if re.match(r"^[（(][一二三四五六七八九十]+[)）]", text):
            return 3
        return 0

    def _extract_headings(self, doc: Document, source: str) -> list[HeadingInfo]:
        result = []
        for idx, para in enumerate(doc.paragraphs):
            lvl = self._detect_heading_level(para)
            if lvl:
                result.append(HeadingInfo(lvl, para.text.strip(), source, idx))
        return result

    def _normalize_style(self, para, level: int):
        font_cn, size_pt = self.cfg.body_font_cn, self.cfg.body_font_size_pt
        if level == 1:
            font_cn, size_pt = self.cfg.heading1_font_cn, self.cfg.heading1_font_size_pt
        elif level == 2:
            font_cn, size_pt = self.cfg.heading2_font_cn, self.cfg.heading2_font_size_pt
        elif level == 3:
            font_cn, size_pt = self.cfg.heading3_font_cn, self.cfg.heading3_font_size_pt
        for run in para.runs:
            if run.text.strip():
                run.font.size = Pt(size_pt)
                run.font.name = font_cn
                rPr = run._element.get_or_add_rPr()
                rFonts = rPr.find(qn("w:rFonts"))
                if rFonts is None:
                    rFonts = OxmlElement("w:rFonts")
                    rPr.append(rFonts)
                rFonts.set(qn("w:eastAsia"), font_cn)
                rFonts.set(qn("w:ascii"), font_cn)
                rFonts.set(qn("w:hAnsi"), font_cn)
                color_el = rPr.find(qn("w:color"))
                if color_el is not None:
                    rPr.remove(color_el)
                highlight = rPr.find(qn("w:highlight"))
                if highlight is not None:
                    rPr.remove(highlight)
        return 1

    def _copy_paragraph(self, src_para, target_doc: Document, level: int):
        new_para = target_doc.add_paragraph()
        if level:
            new_para.style = target_doc.styles[f"Heading {min(level,3)}"]
        pf = new_para.paragraph_format
        if not level:
            from docx.enum.text import WD_LINE_SPACING
            from docx.shared import Pt as _Pt
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf.line_spacing = _Pt(self.cfg.body_line_spacing_lines)
        for run in src_para.runs:
            new_run = new_para.add_run(run.text)
            new_run.bold = run.bold
            new_run.italic = run.italic
            new_run.underline = run.underline
        if level:
            self._normalize_style(new_para, level)
        return new_para

    def _add_toc(self, doc: Document, headings: list[HeadingInfo]):
        toc_para = doc.add_paragraph()
        toc_para.style = doc.styles["Heading 1"]
        run = toc_para.add_run("目  录")
        self._normalize_style(toc_para, 1)
        for h in headings:
            indent = "  " * (h.level - 1)
            p = doc.add_paragraph()
            p.add_run(f"{indent}{h.text}")
            for r in p.runs:
                r.font.size = Pt(self.cfg.body_font_size_pt)
        doc.add_page_break()

    def merge(self, file_paths: list[str], output_path: str,
              with_toc: bool = True) -> MergeResult:
        if not file_paths:
            return MergeResult(False, output_path, 0, error="未提供待合并文件")
        try:
            merged = Document()
            from docx.enum.text import WD_LINE_SPACING
            from docx.shared import Pt as _Pt, Cm
            sec = merged.sections[0]
            sec.page_width = Cm(self.cfg.page_width_cm)
            sec.page_height = Cm(self.cfg.page_height_cm)
            sec.top_margin = Cm(self.cfg.margin_top_cm)
            sec.bottom_margin = Cm(self.cfg.margin_bottom_cm)
            sec.left_margin = Cm(self.cfg.margin_left_cm)
            sec.right_margin = Cm(self.cfg.margin_right_cm)

            all_headings: list[HeadingInfo] = []
            all_content: list[tuple[str, int, object]] = []
            conflicts = 0

            for fp in file_paths:
                p = Path(fp)
                if not p.exists():
                    continue
                src_doc = Document(str(p))
                src_name = p.name
                for para in src_doc.paragraphs:
                    lvl = self._detect_heading_level(para)
                    if lvl:
                        all_headings.append(HeadingInfo(lvl, para.text.strip(), src_name, len(all_content)))
                    all_content.append((src_name, lvl, para))
                    conflicts += self._count_style_conflicts(para)

            if with_toc and all_headings:
                self._add_toc(merged, all_headings)

            for src_name, lvl, para in all_content:
                if not para.text.strip() and lvl == 0:
                    continue
                self._copy_paragraph(para, merged, lvl)

            out = Path(output_path)
            out.parent.mkdir(parents=True, exist_ok=True)
            merged.save(str(out))
            return MergeResult(
                ok=True,
                output_path=str(out),
                heading_count=len(all_headings),
                sources=[Path(f).name for f in file_paths],
                toc=all_headings,
                style_conflicts_resolved=conflicts,
            )
        except Exception as e:
            return MergeResult(False, output_path, 0, error=str(e))

    def _count_style_conflicts(self, para) -> int:
        count = 0
        for run in para.runs:
            if not run.text.strip():
                continue
            rPr = run._element.find(qn("w:rPr"))
            if rPr is not None:
                if rPr.find(qn("w:highlight")) is not None:
                    count += 1
                color = rPr.find(qn("w:color"))
                if color is not None and (color.get(qn("w:val")) or "").lower() not in ("000000", "auto", ""):
                    count += 1
        return count


merger = DocMerger()
